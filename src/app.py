# # src/app.py
# import os
# import json
# import time
# import hashlib
# import secrets
# from typing import Dict, Any
# from flask import Flask, request, jsonify, send_file, abort, send_from_directory
# from datetime import datetime, timedelta
# from io import BytesIO
# from reportlab.lib.pagesizes import A4
# from reportlab.pdfgen import canvas
# from docx import Document

# APP_ROOT = os.path.dirname(os.path.abspath(__file__))
# DATA_DIR = os.path.join(APP_ROOT, "..", "data")
# os.makedirs(DATA_DIR, exist_ok=True)
# QUESTIONS_FILE = os.path.join(DATA_DIR, "questions.json")
# TEMPLATES_FILE = os.path.join(DATA_DIR, "templates.json")

# # Default data for first-run
# DEFAULT_QUESTIONS = [
#     {
#         "id": "q1",
#         "category": "Basics",
#         "label": "Company name",
#         "key": "company_name",
#         "type": "text",
#         "required": True
#     },
#     {
#         "id": "q2",
#         "category": "Contact",
#         "label": "Contact email",
#         "key": "contact_email",
#         "type": "email",
#         "required": False
#     },
#     {
#         "id": "q3",
#         "category": "Data",
#         "label": "Do you collect personal data?",
#         "key": "collects_personal_data",
#         "type": "boolean",
#         "required": True
#     },
#     {
#         "id": "q4",
#         "category": "Tracking",
#         "label": "Do you use cookies?",
#         "key": "uses_cookies",
#         "type": "boolean",
#         "required": True
#     },
#     {
#         "id": "q5",
#         "category": "Legal",
#         "label": "Jurisdiction (country)",
#         "key": "jurisdiction",
#         "type": "text",
#         "required": False
#     }
# ]

# DEFAULT_TEMPLATES = [
#     {
#         "id": "default",
#         "name": "Minimal template",
#         "clauses": {
#             "collects_personal_data_true": "We collect personal information such as name and email required to provide services.",
#             "collects_personal_data_false": "We do not collect personal data beyond anonymous metrics.",
#             "uses_cookies_true": "We use cookies and similar technologies to improve the experience.",
#             "uses_cookies_false": "We do not use cookies.",
#             "default_footer": "By using our service you accept this policy."
#         }
#     }
# ]

# # In-memory session store: token -> {user, expires_at}
# SESSIONS: Dict[str, Dict[str, Any]] = {}

# # Simple cache for previews: key -> (value, expiry)
# CACHE: Dict[str, Dict[str, Any]] = {}

# CACHE_TTL = 10  # seconds for preview caching (fast response requirement)

# app = Flask(__name__, static_folder=os.path.join(APP_ROOT, "..", "static"), static_url_path='/static')

# def load_or_init(file_path: str, default):
#     if not os.path.exists(file_path):
#         with open(file_path, "w", encoding="utf-8") as f:
#             json.dump(default, f, indent=2)
#     with open(file_path, "r", encoding="utf-8") as f:
#         return json.load(f)

# def write_json(file_path: str, data):
#     with open(file_path, "w", encoding="utf-8") as f:
#         json.dump(data, f, indent=2)

# # Initialize data files
# _questions = load_or_init(QUESTIONS_FILE, DEFAULT_QUESTIONS)
# _templates = load_or_init(TEMPLATES_FILE, DEFAULT_TEMPLATES)

# # -------------------------
# # Authentication Utilities
# # -------------------------
# def create_token(user: str, ttl_seconds: int = 900) -> str:
#     token = secrets.token_urlsafe(32)
#     expires_at = time.time() + ttl_seconds
#     SESSIONS[token] = {"user": user, "expires_at": expires_at}
#     return token

# def validate_token(token: str) -> bool:
#     if not token:
#         return False
#     s = SESSIONS.get(token)
#     if not s:
#         return False
#     if time.time() > s["expires_at"]:
#         # auto-logout: remove expired token
#         del SESSIONS[token]
#         return False
#     return True

# def require_auth(fn):
#     def wrapper(*args, **kwargs):
#         token = request.headers.get("Authorization", "")
#         if token.startswith("Bearer "):
#             token = token.split(" ", 1)[1]
#         if not validate_token(token):
#             return jsonify({"error": "unauthorized or session expired"}), 401
#         return fn(*args, **kwargs)
#     wrapper.__name__ = fn.__name__
#     return wrapper

# # -------------------------
# # Helper: clause logic & generation (IM-6, IM-7)
# # -------------------------
# def _select_clauses(answers: dict, template: dict):
#     clauses = []
#     tclauses = template.get("clauses", {})
#     # collects_personal_data
#     if answers.get("collects_personal_data", True):
#         if "collects_personal_data_true" in tclauses:
#             clauses.append(tclauses["collects_personal_data_true"])
#     else:
#         if "collects_personal_data_false" in tclauses:
#             clauses.append(tclauses["collects_personal_data_false"])

#     # cookies
#     if answers.get("uses_cookies", True):
#         if "uses_cookies_true" in tclauses:
#             clauses.append(tclauses["uses_cookies_true"])
#     else:
#         if "uses_cookies_false" in tclauses:
#             clauses.append(tclauses["uses_cookies_false"])

#     # custom additional clause keys
#     for k, v in tclauses.items():
#         if k.startswith("if_"):
#             # key format: if_{answer_key}_{value} => include if
#             try:
#                 _, question_key, expected = k.split("_", 2)
#                 if str(answers.get(question_key, "")).lower() == expected.lower():
#                     clauses.append(v)
#             except Exception:
#                 pass

#     # footer
#     footer = tclauses.get("default_footer", "")
#     return clauses, footer

# def generate_policy_text(answers: dict, template_id: str = "default") -> str:
#     templates = load_or_init(TEMPLATES_FILE, DEFAULT_TEMPLATES)
#     tpl = next((t for t in templates if t.get("id") == template_id), templates[0])
#     company_name = answers.get("company_name", "Company")
#     contact = answers.get("contact_email")
#     jurisdiction = answers.get("jurisdiction", "Unknown")
#     clauses, footer = _select_clauses(answers, tpl)
#     lines = []
#     lines.append(f"Privacy Policy for {company_name}")
#     lines.append(f"Effective date: {datetime.utcnow().date().isoformat()}")
#     lines.append("")
#     lines.append(f"{company_name} (\"we\") values your privacy.")
#     lines.append("")
#     lines.extend(clauses)
#     lines.append("")
#     lines.append(f"We operate in {jurisdiction}.")
#     if contact:
#         lines.append(f"For questions, contact us at {contact}.")
#     lines.append("")
#     if footer:
#         lines.append(footer)
#     return "\n".join(lines)

# # -------------------------
# # Caching helper (fast preview)
# # -------------------------
# def cache_get(key: str):
#     entry = CACHE.get(key)
#     if not entry:
#         return None
#     if entry["expires_at"] < time.time():
#         del CACHE[key]
#         return None
#     return entry["value"]

# def cache_set(key: str, value, ttl=CACHE_TTL):
#     CACHE[key] = {"value": value, "expires_at": time.time() + ttl}

# def make_cache_key(payload: dict) -> str:
#     h = hashlib.sha256(json.dumps(payload, sort_keys=True).encode("utf-8")).hexdigest()
#     return h

# # -------------------------
# # API Routes
# # -------------------------
# @app.route("/api/questions", methods=["GET"])
# def get_questions():
#     qs = load_or_init(QUESTIONS_FILE, DEFAULT_QUESTIONS)
#     # group by category for categorized questions (IM-2)
#     categorized: Dict[str, list] = {}
#     for q in qs:
#         categorized.setdefault(q.get("category", "General"), []).append(q)
#     return jsonify({"questions": qs, "categorized": categorized})

# @app.route("/api/questions", methods=["POST"])
# @require_auth
# def add_question():
#     body = request.get_json(force=True, silent=True) or {}
#     qs = load_or_init(QUESTIONS_FILE, DEFAULT_QUESTIONS)
#     # validate minimal fields
#     if not body.get("id") or not body.get("label") or not body.get("key"):
#         return jsonify({"error": "missing id/label/key"}), 400
#     qs.append(body)
#     write_json(QUESTIONS_FILE, qs)
#     return jsonify({"ok": True, "question": body}), 201

# @app.route("/api/questions/<qid>", methods=["PUT", "DELETE"])
# @require_auth
# def modify_question(qid):
#     qs = load_or_init(QUESTIONS_FILE, DEFAULT_QUESTIONS)
#     found = next((q for q in qs if q.get("id") == qid), None)
#     if not found:
#         return jsonify({"error": "not found"}), 404
#     if request.method == "DELETE":
#         qs = [q for q in qs if q.get("id") != qid]
#         write_json(QUESTIONS_FILE, qs)
#         return jsonify({"ok": True}), 200
#     body = request.get_json(force=True, silent=True) or {}
#     # update allowed fields
#     for k in ("label", "category", "type", "key", "required"):
#         if k in body:
#             found[k] = body[k]
#     write_json(QUESTIONS_FILE, qs)
#     return jsonify({"ok": True, "question": found}), 200

# # Templates management (IM-16)
# @app.route("/api/templates", methods=["GET"])
# def get_templates():
#     t = load_or_init(TEMPLATES_FILE, DEFAULT_TEMPLATES)
#     return jsonify({"templates": t})

# @app.route("/api/templates", methods=["POST"])
# @require_auth
# def add_template():
#     body = request.get_json(force=True, silent=True) or {}
#     if not body.get("id") or not body.get("name"):
#         return jsonify({"error": "missing id or name"}), 400
#     t = load_or_init(TEMPLATES_FILE, DEFAULT_TEMPLATES)
#     t.append(body)
#     write_json(TEMPLATES_FILE, t)
#     return jsonify({"ok": True, "template": body}), 201

# @app.route("/api/templates/<tid>", methods=["PUT", "DELETE"])
# @require_auth
# def modify_template(tid):
#     t = load_or_init(TEMPLATES_FILE, DEFAULT_TEMPLATES)
#     found = next((x for x in t if x.get("id") == tid), None)
#     if not found:
#         return jsonify({"error": "not found"}), 404
#     if request.method == "DELETE":
#         t = [x for x in t if x.get("id") != tid]
#         write_json(TEMPLATES_FILE, t)
#         return jsonify({"ok": True}), 200
#     body = request.get_json(force=True, silent=True) or {}
#     found.update(body)
#     write_json(TEMPLATES_FILE, t)
#     return jsonify({"ok": True, "template": found}), 200

# # Login / session (IM-19, IM-20)
# @app.route("/api/login", methods=["POST"])
# def login():
#     body = request.get_json(force=True, silent=True) or {}
#     username = body.get("username")
#     password = body.get("password")
#     # Minimal auth for demo: credentials in env or default admin/admin
#     admin_user = os.environ.get("PPG_ADMIN_USER", "admin")
#     admin_pass = os.environ.get("PPG_ADMIN_PASS", "admin")
#     if username == admin_user and password == admin_pass:
#         token = create_token(username, ttl_seconds=900)  # 15 minutes default, auto-logout ok
#         return jsonify({"token": token, "expires_in": 900}), 200
#     return jsonify({"error": "invalid credentials"}), 401

# # Generate policy and preview endpoints
# @app.route("/api/generate", methods=["POST"])
# def generate():
#     payload = request.get_json(force=True, silent=True) or {}
#     # produce text policy
#     try:
#         template_id = payload.get("template_id", "default")
#         answers = payload.get("answers", {})
#         policy = generate_policy_text(answers, template_id)
#     except Exception as e:
#         return jsonify({"error": str(e)}), 400
#     return jsonify({"policy": policy}), 200

# @app.route("/api/preview", methods=["POST"])
# def preview():
#     payload = request.get_json(force=True, silent=True) or {}
#     key = make_cache_key(payload)
#     cached = cache_get(key)
#     if cached:
#         return jsonify({"policy": cached, "cached": True})
#     template_id = payload.get("template_id", "default")
#     answers = payload.get("answers", {})
#     policy = generate_policy_text(answers, template_id)
#     cache_set(key, policy)
#     return jsonify({"policy": policy, "cached": False})

# # Export as PDF and DOCX (IM-12 & IM-13)
# @app.route("/api/export/pdf", methods=["POST"])
# def export_pdf():
#     payload = request.get_json(force=True, silent=True) or {}
#     answers = payload.get("answers", {})
#     template_id = payload.get("template_id", "default")
#     txt = generate_policy_text(answers, template_id)
#     # Create PDF in memory
#     buf = BytesIO()
#     c = canvas.Canvas(buf, pagesize=A4)
#     width, height = A4
#     margin = 50
#     y = height - margin
#     lines = txt.split("\n")
#     c.setFont("Helvetica", 11)
#     for line in lines:
#         if y < margin:
#             c.showPage()
#             y = height - margin
#             c.setFont("Helvetica", 11)
#         c.drawString(margin, y, line)
#         y -= 14
#     c.save()
#     buf.seek(0)
#     return send_file(buf, mimetype="application/pdf", as_attachment=True, download_name="privacy_policy.pdf")

# @app.route("/api/export/docx", methods=["POST"])
# def export_docx():
#     payload = request.get_json(force=True, silent=True) or {}
#     answers = payload.get("answers", {})
#     template_id = payload.get("template_id", "default")
#     txt = generate_policy_text(answers, template_id)
#     doc = Document()
#     for line in txt.split("\n"):
#         doc.add_paragraph(line)
#     buf = BytesIO()
#     doc.save(buf)
#     buf.seek(0)
#     return send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
#                      as_attachment=True, download_name="privacy_policy.docx")

# # Serve static UI (IM-23 responsive UI)
# @app.route("/", methods=["GET"])
# def index():
#     # Serve the static single-page app
#     return send_from_directory(os.path.join(APP_ROOT, "..", "static"), "index.html")

# # Health check + fast response tip (IM-22)
# @app.route("/health", methods=["GET"])
# def health():
#     return jsonify({"status": "ok", "time": datetime.utcnow().isoformat()})

# # Run
# if __name__ == "__main__":
#     # For dev only: debug True. Do NOT enable debug in production.
#     app.run(debug=True, host="0.0.0.0", port=5000)



# # #testing for categorical questions :

# # # src/app.py
# # from flask import Flask, jsonify

# # app = Flask(__name__)

# # @app.route('/api/questions')
# # def get_questions():
# #     categorized_questions = {
# #         "Personal Info": [
# #             {"id": "q1", "label": "What is your name?", "type": "text"},
# #             {"id": "q2", "label": "Email address", "type": "email"}
# #         ],
# #         "Privacy Preferences": [
# #             {"id": "q3", "label": "Do you accept cookies?", "type": "boolean"}
# #         ]
# #     }
# #     return jsonify({"categorized": categorized_questions})

# # if __name__ == "__main__":
# #     app.run(debug=True)








# from flask import Flask, jsonify, request, send_file
# from datetime import datetime, timedelta, UTC
# from pymongo import MongoClient
# from dotenv import load_dotenv
# import os, io
# from reportlab.pdfgen import canvas
# from docx import Document

# app = Flask(__name__, static_folder="../static", static_url_path="/")

# # ------------------------------
# # MongoDB Setup
# # ------------------------------
# load_dotenv()
# client = MongoClient(os.getenv("MONGO_URI"))
# db = client["privacy_policy_db"]
# questions_col = db["questions"]
# templates_col = db["templates"]

# # Admin token management
# ADMIN_TOKEN = None
# TOKEN_EXPIRY = None


# # ------------------------------
# # Frontend Route
# # ------------------------------
# @app.route("/")
# def home():
#     return app.send_static_file("index.html")


# # ------------------------------
# # API Routes
# # ------------------------------

# @app.route("/api/questions", methods=["GET"])
# def get_questions():
#     """Fetch categorized questions from MongoDB."""
#     questions = list(questions_col.find({}, {"_id": 0}))
#     categorized = {}
#     for q in questions:
#         categorized.setdefault(q["category"], []).append(q)
#     return jsonify({"questions": questions, "categorized": categorized})


# @app.route("/api/questions", methods=["POST"])
# def add_question():
#     """Add a new question (admin only)."""
#     token = request.headers.get("Authorization")
#     if not valid_token(token):
#         return jsonify({"error": "Unauthorized"}), 401
#     data = request.get_json()
#     if not data.get("category") or not data.get("question"):
#         return jsonify({"error": "Missing fields"}), 400
#     questions_col.insert_one(data)
#     return jsonify({"message": "Question added"}), 201


# @app.route("/api/questions/<string:key>", methods=["DELETE"])
# def delete_question(key):
#     """Delete question by key (admin only)."""
#     token = request.headers.get("Authorization")
#     if not valid_token(token):
#         return jsonify({"error": "Unauthorized"}), 401
#     result = questions_col.delete_one({"key": key})
#     if result.deleted_count == 0:
#         return jsonify({"error": "Not found"}), 404
#     return jsonify({"message": "Question deleted"}), 200


# @app.route("/api/templates", methods=["GET"])
# def get_templates():
#     templates = list(templates_col.find({}, {"_id": 0}))
#     return jsonify({"templates": templates})


# @app.route("/api/templates", methods=["POST"])
# def add_template():
#     """Add a new template (admin only)."""
#     token = request.headers.get("Authorization")
#     if not valid_token(token):
#         return jsonify({"error": "Unauthorized"}), 401
#     data = request.get_json()
#     if not data.get("title") or not data.get("content"):
#         return jsonify({"error": "Missing fields"}), 400
#     templates_col.insert_one(data)
#     return jsonify({"message": "Template added"}), 201


# @app.route("/api/templates/<string:title>", methods=["DELETE"])
# def delete_template(title):
#     """Delete a template by title."""
#     token = request.headers.get("Authorization")
#     if not valid_token(token):
#         return jsonify({"error": "Unauthorized"}), 401
#     result = templates_col.delete_one({"title": title})
#     if result.deleted_count == 0:
#         return jsonify({"error": "Not found"}), 404
#     return jsonify({"message": "Template deleted"}), 200


# @app.route("/api/login", methods=["POST"])
# def login():
#     """Admin login."""
#     creds = request.get_json()
#     if creds["username"] == "admin" and creds["password"] == "admin":
#         global ADMIN_TOKEN, TOKEN_EXPIRY
#         ADMIN_TOKEN = "valid_token"
#         TOKEN_EXPIRY = datetime.now(UTC) + timedelta(seconds=900)
#         return jsonify({"message": "Login success", "token": ADMIN_TOKEN, "expiry": 900})
#     return jsonify({"error": "Invalid credentials"}), 403


# def valid_token(token):
#     global ADMIN_TOKEN, TOKEN_EXPIRY
#     return token == ADMIN_TOKEN and TOKEN_EXPIRY and TOKEN_EXPIRY > datetime.now(UTC)


# # ------------------------------
# # Generate Privacy Policy Text
# # ------------------------------
# def generate_policy_text(answers, template_title="Default Template"):
#     tpl = templates_col.find_one({"title": template_title})
#     if not tpl:
#         return "Default Privacy Policy template not found."

#     text = tpl["content"]
#     replacements = {
#         "{company_name}": answers.get("company_name", "Your Company"),
#         "{contact_email}": answers.get("contact_email", "contact@example.com"),
#         "{cookies_text}": "use cookies" if answers.get("uses_cookies") else "do not use cookies",
#         "{jurisdiction}": answers.get("jurisdiction", "your jurisdiction"),
#         "{date}": datetime.now(UTC).date().isoformat()
#     }

#     for key, val in replacements.items():
#         text = text.replace(key, str(val))

#     return text


# @app.route("/api/preview", methods=["POST"])
# def preview_policy():
#     """Generate live preview of privacy policy."""
#     data = request.get_json()
#     answers = data.get("answers", {})
#     policy_text = generate_policy_text(answers)
#     return jsonify({"preview": policy_text})


# @app.route("/api/export/pdf", methods=["POST"])
# def export_pdf():
#     """Export privacy policy as PDF."""
#     data = request.get_json()
#     buffer = io.BytesIO()
#     p = canvas.Canvas(buffer)
#     y = 800
#     for line in data.get("content", "").splitlines():
#         p.drawString(100, y, line)
#         y -= 15
#     p.save()
#     buffer.seek(0)
#     return send_file(buffer, as_attachment=True, download_name="policy.pdf", mimetype="application/pdf")


# @app.route("/api/export/docx", methods=["POST"])
# def export_docx():
#     """Export privacy policy as DOCX."""
#     data = request.get_json()
#     doc = Document()
#     doc.add_heading("Privacy Policy", 0)
#     doc.add_paragraph(data.get("content", ""))
#     file_stream = io.BytesIO()
#     doc.save(file_stream)
#     file_stream.seek(0)
#     return send_file(file_stream, as_attachment=True, download_name="policy.docx")


# @app.route("/api/health")
# def health():
#     """CI/CD health check route."""
#     return jsonify({"status": "ok", "time": datetime.now(UTC).isoformat()})


# if __name__ == "__main__":
#     app.run(debug=True)






"""
Privacy Policy Generator - Flask Backend
Complete implementation of all user stories with CI/CD compliance

Story Mapping:
- IM-2: Categorized Questions (GET /api/questions)
- IM-3: Form Navigation (handled by frontend, backend supports it)
- IM-6: Insert Details / Template Selection (GET /api/templates)
- IM-7: Clause Logic (generate_policy_text function)
- IM-9: Live Preview (POST /api/preview)
- IM-10: Inline Edit (handled by frontend)
- IM-12: PDF Export (POST /api/export/pdf)
- IM-13: DOCX Export (POST /api/export/docx)
- IM-15: Manage Questions (POST/PUT/DELETE /api/questions)
- IM-16: Manage Templates (POST/PUT/DELETE /api/templates)
- IM-18: HTTPS Security (security headers)
- IM-19: User Login (POST /api/login)
- IM-20: Auto Logout (session expiry logic)
- IM-22: Fast Response (database indexing, caching)
- IM-23: Responsive UI (frontend handles, backend supports)
"""

from flask import Flask, jsonify, request, send_file
from datetime import datetime, timedelta, UTC
from pymongo import MongoClient
from dotenv import load_dotenv
import os
import io
import secrets
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.units import inch
from docx import Document
from docx.shared import Pt, RGBColor
from werkzeug.security import generate_password_hash, check_password_hash

# ==============================
# Flask App Configuration
# ==============================
app = Flask(__name__, static_folder="../static", static_url_path="/")
app.config['SECRET_KEY'] = os.getenv('SECRET_KEY', 'dev-secret-key-change-in-production')
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16MB max request size

# ==============================
# MongoDB Setup
# ==============================
load_dotenv()

# Get MongoDB URI from environment
MONGO_URI = os.getenv("MONGO_URI")
if not MONGO_URI:
    raise ValueError("MONGO_URI not found in environment variables. Please check your .env file.")

# Connect to MongoDB
try:
    client = MongoClient(MONGO_URI, serverSelectionTimeoutMS=5000)
    # Test connection
    client.admin.command('ping')
    print("✅ Successfully connected to MongoDB")
except Exception as e:
    print(f"❌ Failed to connect to MongoDB: {e}")
    raise

# Database and Collections
db = client["privacy_policy_db"]
questions_col = db["questions"]
templates_col = db["templates"]
users_col = db["users"]
sessions_col = db["sessions"]

# Create indexes for performance (IM-22: Fast Response)
questions_col.create_index("key", unique=True)
questions_col.create_index("category")
templates_col.create_index("title", unique=True)
sessions_col.create_index("token")
sessions_col.create_index("expiry", expireAfterSeconds=0)  # TTL index

# ==============================
# Security Headers (IM-18: HTTPS Security)
# ==============================
@app.after_request
def add_security_headers(response):
    """Add security headers to all responses."""
    # Enforce HTTPS in production
    response.headers['Strict-Transport-Security'] = 'max-age=31536000; includeSubDomains'
    # Prevent clickjacking
    response.headers['X-Frame-Options'] = 'SAMEORIGIN'
    # Prevent MIME sniffing
    response.headers['X-Content-Type-Options'] = 'nosniff'
    # XSS Protection
    response.headers['X-XSS-Protection'] = '1; mode=block'
    # Content Security Policy
    response.headers['Content-Security-Policy'] = "default-src 'self' https://cdn.jsdelivr.net https://cdnjs.cloudflare.com; script-src 'self' 'unsafe-inline' https://cdn.jsdelivr.net; style-src 'self' 'unsafe-inline' https://cdn.jsdelivr.net https://cdnjs.cloudflare.com;"
    # Referrer Policy
    response.headers['Referrer-Policy'] = 'strict-origin-when-cross-origin'
    # CORS headers (if needed)
    response.headers['Access-Control-Allow-Origin'] = '*'
    response.headers['Access-Control-Allow-Methods'] = 'GET, POST, PUT, DELETE, OPTIONS'
    response.headers['Access-Control-Allow-Headers'] = 'Content-Type, Authorization'
    
    return response


# ==============================
# Session Management (In-Memory Cache)
# ==============================
ADMIN_SESSIONS = {}  # In-memory cache for active sessions


def cleanup_expired_sessions():
    """Remove expired sessions from memory cache."""
    current_time = datetime.now(UTC)
    expired_tokens = [token for token, session in ADMIN_SESSIONS.items() 
                     if session['expiry'] < current_time]
    for token in expired_tokens:
        del ADMIN_SESSIONS[token]


# ==============================
# Initialize Default Admin User (IM-19: User Login)
# ==============================
def init_admin_user():
    """Initialize default admin user if not exists."""
    try:
        if users_col.count_documents({"username": "admin"}) == 0:
            users_col.insert_one({
                "username": "admin",
                "password": generate_password_hash("admin"),
                "role": "admin",
                "created_at": datetime.now(UTC)
            })
            print("✅ Default admin user created (username: admin, password: admin)")
    except Exception as e:
        print(f"⚠️  Warning: Could not create default admin user: {e}")

init_admin_user()


# ==============================
# Frontend Route
# ==============================
@app.route("/")
def home():
    """Serve the main HTML page."""
    return app.send_static_file("index.html")


# ==============================
# Authentication Helpers
# ==============================
def valid_token(token):
    """
    Check if authentication token is valid.
    IM-19: User Login
    IM-20: Auto Logout (15-minute expiry)
    """
    if not token:
        return False
    
    # Cleanup expired sessions
    cleanup_expired_sessions()
    
    # Check in-memory cache first (IM-22: Fast Response)
    session = ADMIN_SESSIONS.get(token)
    if session and session["expiry"] > datetime.now(UTC):
        return True
    
    # Check database as fallback
    try:
        session = sessions_col.find_one({"token": token})
        if session and session["expiry"] > datetime.now(UTC):
            # Cache it for faster access
            ADMIN_SESSIONS[token] = {
                "username": session["username"],
                "expiry": session["expiry"]
            }
            return True
    except Exception as e:
        print(f"Error checking token: {e}")
    
    # Clean up invalid/expired session
    if token in ADMIN_SESSIONS:
        del ADMIN_SESSIONS[token]
    try:
        sessions_col.delete_one({"token": token})
    except:
        pass
    
    return False


# ==============================
# Authentication Routes (IM-19: User Login, IM-20: Auto Logout)
# ==============================
@app.route("/api/login", methods=["POST"])
def login():
    """
    Admin login endpoint.
    IM-19: User Login
    IM-20: Auto Logout (creates session with 15-minute expiry)
    """
    try:
        creds = request.get_json()
        
        # Validate input
        if not creds or not creds.get("username") or not creds.get("password"):
            return jsonify({"error": "Missing username or password"}), 400
        
        # Find user in database
        user = users_col.find_one({"username": creds["username"]})
        
        # Verify credentials
        if not user or not check_password_hash(user["password"], creds["password"]):
            return jsonify({"error": "Invalid credentials"}), 403
        
        # Generate secure session token
        token = secrets.token_urlsafe(32)
        expiry = datetime.now(UTC) + timedelta(seconds=900)  # 15 minutes (IM-20)
        
        # Store session in database
        session_data = {
            "token": token,
            "username": user["username"],
            "role": user.get("role", "admin"),
            "expiry": expiry,
            "created_at": datetime.now(UTC)
        }
        sessions_col.insert_one(session_data)
        
        # Cache session in memory for fast access (IM-22)
        ADMIN_SESSIONS[token] = {
            "username": user["username"],
            "expiry": expiry
        }
        
        return jsonify({
            "message": "Login successful",
            "token": token,
            "expiry": 900,  # seconds
            "username": user["username"]
        }), 200
        
    except Exception as e:
        print(f"Login error: {e}")
        return jsonify({"error": "Login failed"}), 500


@app.route("/api/logout", methods=["POST"])
def logout():
    """
    Logout endpoint - invalidates session.
    IM-19: User Login
    """
    try:
        token = request.headers.get("Authorization")
        if token:
            # Remove from memory cache
            if token in ADMIN_SESSIONS:
                del ADMIN_SESSIONS[token]
            # Remove from database
            sessions_col.delete_one({"token": token})
        
        return jsonify({"message": "Logged out successfully"}), 200
    except Exception as e:
        print(f"Logout error: {e}")
        return jsonify({"error": "Logout failed"}), 500


# ==============================
# Questions API Routes (IM-2: Categorized Questions, IM-15: Manage Questions)
# ==============================
@app.route("/api/questions", methods=["GET"])
def get_questions():
    """
    Fetch all questions, categorized.
    IM-2: Categorized Questions
    IM-22: Fast Response (uses database indexing)
    """
    try:
        # Fetch questions sorted by order
        questions = list(questions_col.find({}, {"_id": 0}).sort("order", 1))
        
        # Categorize questions
        categorized = {}
        for q in questions:
            category = q.get("category", "General")
            if category not in categorized:
                categorized[category] = []
            categorized[category].append(q)
        
        return jsonify({
            "questions": questions,
            "categorized": categorized
        }), 200
        
    except Exception as e:
        print(f"Error fetching questions: {e}")
        return jsonify({"error": "Failed to fetch questions"}), 500


@app.route("/api/questions", methods=["POST"])
def add_question():
    """
    Add a new question (admin only).
    IM-15: Manage Questions
    """
    # Check authentication
    token = request.headers.get("Authorization")
    if not valid_token(token):
        return jsonify({"error": "Unauthorized"}), 401
    
    try:
        data = request.get_json()
        
        # Validate required fields
        required_fields = ["key", "category", "question", "type"]
        if not all(field in data for field in required_fields):
            return jsonify({"error": "Missing required fields: key, category, question, type"}), 400
        
        # Check if key already exists
        if questions_col.find_one({"key": data["key"]}):
            return jsonify({"error": "Question with this key already exists"}), 409
        
        # Prepare question document
        question_data = {
            "key": data["key"],
            "category": data["category"],
            "question": data["question"],
            "type": data["type"],
            "required": data.get("required", False),
            "order": data.get("order", 999),
            "options": data.get("options", []),
            "dependsOn": data.get("dependsOn", None),
            "created_at": datetime.now(UTC)
        }
        
        # Insert into database
        questions_col.insert_one(question_data)
        
        return jsonify({"message": "Question added successfully"}), 201
        
    except Exception as e:
        print(f"Error adding question: {e}")
        return jsonify({"error": "Failed to add question"}), 500


@app.route("/api/questions/<string:key>", methods=["PUT"])
def update_question(key):
    """
    Update an existing question (admin only).
    IM-15: Manage Questions
    """
    # Check authentication
    token = request.headers.get("Authorization")
    if not valid_token(token):
        return jsonify({"error": "Unauthorized"}), 401
    
    try:
        data = request.get_json()
        
        # Prepare update document
        update_data = {
            "category": data.get("category"),
            "question": data.get("question"),
            "type": data.get("type"),
            "required": data.get("required", False),
            "order": data.get("order", 999),
            "options": data.get("options", []),
            "dependsOn": data.get("dependsOn"),
            "updated_at": datetime.now(UTC)
        }
        
        # Update in database
        result = questions_col.update_one({"key": key}, {"$set": update_data})
        
        if result.matched_count == 0:
            return jsonify({"error": "Question not found"}), 404
        
        return jsonify({"message": "Question updated successfully"}), 200
        
    except Exception as e:
        print(f"Error updating question: {e}")
        return jsonify({"error": "Failed to update question"}), 500


@app.route("/api/questions/<string:key>", methods=["DELETE"])
def delete_question(key):
    """
    Delete a question by key (admin only).
    IM-15: Manage Questions
    """
    # Check authentication
    token = request.headers.get("Authorization")
    if not valid_token(token):
        return jsonify({"error": "Unauthorized"}), 401
    
    try:
        result = questions_col.delete_one({"key": key})
        
        if result.deleted_count == 0:
            return jsonify({"error": "Question not found"}), 404
        
        return jsonify({"message": "Question deleted successfully"}), 200
        
    except Exception as e:
        print(f"Error deleting question: {e}")
        return jsonify({"error": "Failed to delete question"}), 500


# ==============================
# Templates API Routes (IM-6: Insert Details, IM-16: Manage Templates)
# ==============================
@app.route("/api/templates", methods=["GET"])
def get_templates():
    """
    Get all available templates.
    IM-6: Insert Details (Template Selection)
    IM-22: Fast Response
    """
    try:
        templates = list(templates_col.find({}, {"_id": 0}))
        return jsonify({"templates": templates}), 200
    except Exception as e:
        print(f"Error fetching templates: {e}")
        return jsonify({"error": "Failed to fetch templates"}), 500


@app.route("/api/templates", methods=["POST"])
def add_template():
    """
    Add a new template (admin only).
    IM-16: Manage Templates
    """
    # Check authentication
    token = request.headers.get("Authorization")
    if not valid_token(token):
        return jsonify({"error": "Unauthorized"}), 401
    
    try:
        data = request.get_json()
        
        # Validate required fields
        if not data.get("title") or not data.get("content"):
            return jsonify({"error": "Missing required fields: title, content"}), 400
        
        # Check if template already exists
        if templates_col.find_one({"title": data["title"]}):
            return jsonify({"error": "Template with this title already exists"}), 409
        
        # Prepare template document
        template_data = {
            "title": data["title"],
            "content": data["content"],
            "description": data.get("description", ""),
            "created_at": datetime.now(UTC)
        }
        
        # Insert into database
        templates_col.insert_one(template_data)
        
        return jsonify({"message": "Template added successfully"}), 201
        
    except Exception as e:
        print(f"Error adding template: {e}")
        return jsonify({"error": "Failed to add template"}), 500


@app.route("/api/templates/<string:title>", methods=["PUT"])
def update_template(title):
    """
    Update a template (admin only).
    IM-16: Manage Templates
    """
    # Check authentication
    token = request.headers.get("Authorization")
    if not valid_token(token):
        return jsonify({"error": "Unauthorized"}), 401
    
    try:
        data = request.get_json()
        
        # Prepare update document
        update_data = {
            "content": data.get("content"),
            "description": data.get("description", ""),
            "updated_at": datetime.now(UTC)
        }
        
        # Update in database
        result = templates_col.update_one({"title": title}, {"$set": update_data})
        
        if result.matched_count == 0:
            return jsonify({"error": "Template not found"}), 404
        
        return jsonify({"message": "Template updated successfully"}), 200
        
    except Exception as e:
        print(f"Error updating template: {e}")
        return jsonify({"error": "Failed to update template"}), 500


@app.route("/api/templates/<string:title>", methods=["DELETE"])
def delete_template(title):
    """
    Delete a template by title (admin only).
    IM-16: Manage Templates
    """
    # Check authentication
    token = request.headers.get("Authorization")
    if not valid_token(token):
        return jsonify({"error": "Unauthorized"}), 401
    
    try:
        result = templates_col.delete_one({"title": title})
        
        if result.deleted_count == 0:
            return jsonify({"error": "Template not found"}), 404
        
        return jsonify({"message": "Template deleted successfully"}), 200
        
    except Exception as e:
        print(f"Error deleting template: {e}")
        return jsonify({"error": "Failed to delete template"}), 500


# ==============================
# Policy Generation (IM-7: Clause Logic)
# ==============================
def generate_policy_text(answers, template_title="Default Template"):
    """
    Generate privacy policy text with conditional clauses.
    IM-7: Clause Logic
    IM-9: Live Preview (used by preview endpoint)
    """
    try:
        # Fetch template
        tpl = templates_col.find_one({"title": template_title})
        if not tpl:
            return "Error: Template not found. Please ensure the template exists in the database."
        
        text = tpl["content"]
        
        # Basic replacements
        replacements = {
            "{company_name}": answers.get("company_name", "Your Company"),
            "{contact_email}": answers.get("contact_email", "contact@example.com"),
            "{website_url}": answers.get("website_url", "https://example.com"),
            "{jurisdiction}": answers.get("jurisdiction", "your jurisdiction"),
            "{date}": datetime.now(UTC).strftime("%B %d, %Y"),
            "{data_retention}": answers.get("data_retention", "as required by law")
        }
        
        # IM-7: Conditional Clause Logic
        
        # Cookies clause
        uses_cookies = answers.get("uses_cookies")
        if uses_cookies == "true" or uses_cookies == True or uses_cookies == "Yes":
            replacements["{cookies_text}"] = "use cookies to enhance user experience and analyze website traffic"
        else:
            replacements["{cookies_text}"] = "do not use cookies"
        
        # Data collection clause
        collects_data = answers.get("collects_personal_data")
        if collects_data == "true" or collects_data == True or collects_data == "Yes":
            data_types = answers.get("data_types", "name, email address, and other information you voluntarily provide")
            replacements["{data_collection_text}"] = f"We collect the following types of personal data: {data_types}."
        else:
            replacements["{data_collection_text}"] = "We do not collect personal data from our users."
        
        # Data usage clause
        replacements["{data_usage_text}"] = "We use your data to provide and improve our services, communicate with you regarding your account or our services, ensure security, and comply with legal obligations."
        
        # Third-party sharing clause
        third_party = answers.get("third_party_sharing")
        if third_party == "true" or third_party == True or third_party == "Yes":
            replacements["{third_party_text}"] = "We may share your personal data with trusted third-party service providers who assist us in operating our website, conducting our business, or servicing you. These parties are obligated to keep your information confidential."
        else:
            replacements["{third_party_text}"] = "We do not share your personal data with third parties except as required by law."
        
        # User rights clause
        user_rights = answers.get("user_rights", [])
        if isinstance(user_rights, str):
            user_rights = [user_rights]
        if user_rights and len(user_rights) > 0:
            rights_list = ", ".join(user_rights)
            replacements["{user_rights_text}"] = f"Under applicable privacy laws, you have the following rights regarding your personal data: {rights_list}. To exercise these rights, please contact us using the information provided below."
        else:
            replacements["{user_rights_text}"] = "You have rights regarding your personal data as provided under applicable privacy laws. Please contact us for more information."
        
        # Legal basis for processing (GDPR)
        replacements["{legal_basis_text}"] = "We process your personal data based on one or more of the following legal bases: your consent, performance of a contract with you, compliance with legal obligations, protection of vital interests, public interest, or our legitimate business interests."
        
        # Apply all replacements
        for key, val in replacements.items():
            text = text.replace(key, str(val))
        
        return text
        
    except Exception as e:
        print(f"Error generating policy text: {e}")
        return f"Error generating policy: {str(e)}"


# ==============================
# Preview Route (IM-9: Live Preview)
# ==============================
@app.route("/api/preview", methods=["POST"])
def preview_policy():
    """
    Generate live preview of privacy policy.
    IM-9: Live Preview
    IM-7: Clause Logic (uses generate_policy_text)
    """
    try:
        data = request.get_json()
        
        if not data:
            return jsonify({"error": "No data provided"}), 400
        
        answers = data.get("answers", {})
        template_title = data.get("template", "Default Template")
        
        # Generate policy text
        policy_text = generate_policy_text(answers, template_title)
        
        return jsonify({
            "preview": policy_text,
            "generated_at": datetime.now(UTC).isoformat()
        }), 200
        
    except Exception as e:
        print(f"Error generating preview: {e}")
        return jsonify({"error": "Failed to generate preview"}), 500


# ==============================
# Export Routes (IM-12: PDF Export, IM-13: DOCX Export)
# ==============================
@app.route("/api/export/pdf", methods=["POST"])
def export_pdf():
    """
    Export privacy policy as PDF.
    IM-12: PDF Export
    """
    try:
        data = request.get_json()
        content = data.get("content", "")
        
        if not content:
            return jsonify({"error": "No content provided for export"}), 400
        
        # Create PDF buffer
        buffer = io.BytesIO()
        
        # Create PDF with better formatting using reportlab
        doc = SimpleDocTemplate(buffer, pagesize=letter,
                               topMargin=0.75*inch, bottomMargin=0.75*inch,
                               leftMargin=0.75*inch, rightMargin=0.75*inch)
        
        # Story to hold document elements
        story = []
        styles = getSampleStyleSheet()
        
        # Customize styles
        title_style = styles['Heading1']
        heading_style = styles['Heading2']
        body_style = styles['BodyText']
        body_style.fontSize = 11
        body_style.leading = 14
        
        # Split content into paragraphs
        paragraphs = content.split('\n\n')
        
        for para in paragraphs:
            if para.strip():
                # Detect if it's a heading (all caps or starts with digit)
                if para.strip().isupper() and len(para.strip()) < 100:
                    p = Paragraph(para.strip(), heading_style)
                elif para.strip() and para.strip()[0].isdigit() and '. ' in para[:10]:
                    p = Paragraph(para.strip(), heading_style)
                else:
                    # Regular paragraph
                    para_text = para.replace('\n', '<br/>')
                    p = Paragraph(para_text, body_style)
                
                story.append(p)
                story.append(Spacer(1, 0.15*inch))
        
        # Build PDF
        doc.build(story)
        buffer.seek(0)
        
        return send_file(
            buffer,
            as_attachment=True,
            download_name="privacy_policy.pdf",
            mimetype="application/pdf"
        )
        
    except Exception as e:
        print(f"Error exporting PDF: {e}")
        return jsonify({"error": "Failed to export PDF"}), 500


@app.route("/api/export/docx", methods=["POST"])
def export_docx():
    """
    Export privacy policy as DOCX.
    IM-13: DOCX Export
    """
    try:
        data = request.get_json()
        content = data.get("content", "")
        
        if not content:
            return jsonify({"error": "No content provided for export"}), 400
        
        # Create DOCX document
        doc = Document()
        
        # Add main title
        title = doc.add_heading("Privacy Policy", 0)
        title.alignment = 1  # Center alignment
        
        # Split content into paragraphs
        paragraphs = content.split('\n\n')
        
        for para in paragraphs:
            if para.strip():
                # Detect headings
                if para.strip().isupper() and len(para.strip()) < 100:
                    # Section heading
                    doc.add_heading(para.strip(), level=2)
                elif para.strip() and para.strip()[0].isdigit() and '. ' in para[:10]:
                    # Numbered heading
                    doc.add_heading(para.strip(), level=2)
                else:
                    # Regular paragraph
                    p = doc.add_paragraph(para.strip())
                    # Format paragraph
                    for run in p.runs:
                        run.font.size = Pt(11)
                        run.font.name = 'Calibri'
        
        # Save to BytesIO buffer
        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)
        
        return send_file(
            file_stream,
            as_attachment=True,
            download_name="privacy_policy.docx",
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        
    except Exception as e:
        print(f"Error exporting DOCX: {e}")
        return jsonify({"error": "Failed to export DOCX"}), 500


# ==============================
# Health Check (CI/CD Requirement)
# ==============================
@app.route("/api/health")
def health():
    """
    Health check endpoint for CI/CD pipeline.
    Tests database connectivity and returns status.
    """
    try:
        # Test database connection
        client.admin.command('ping')
        
        return jsonify({
            "status": "ok",
            "time": datetime.now(UTC).isoformat(),
            "database": "connected",
            "version": "1.0.0"
        }), 200
        
    except Exception as e:
        print(f"Health check failed: {e}")
        return jsonify({
            "status": "error",
            "time": datetime.now(UTC).isoformat(),
            "database": "disconnected",
            "error": str(e)
        }), 500


# ==============================
# Error Handlers
# ==============================
@app.errorhandler(404)
def not_found(error):
    """Handle 404 errors."""
    return jsonify({"error": "Resource not found"}), 404


@app.errorhandler(500)
def internal_error(error):
    """Handle 500 errors."""
    return jsonify({"error": "Internal server error"}), 500


@app.errorhandler(405)
def method_not_allowed(error):
    """Handle 405 errors."""
    return jsonify({"error": "Method not allowed"}), 405


# ==============================
# Run Application
# ==============================
if __name__ == "__main__":
    # In production, use a production WSGI server like Gunicorn
    # For development:
    app.run(
        debug=os.getenv("FLASK_DEBUG", "True") == "True",
        host="0.0.0.0",
        port=int(os.getenv("PORT", 5000))
    )





